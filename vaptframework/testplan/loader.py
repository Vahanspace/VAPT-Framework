"""Load the VAPT test plan (xlsx) as structured TestCase records.

The spreadsheet is the human-facing source of truth: analysts read and annotate it, and
the framework reads the same file so automation and the worksheet never drift apart.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

import openpyxl

TEST_CASE_SHEET = "VAPT Test Cases"

# Header label -> TestCase attribute
_COLUMN_MAP = {
    "Test ID": "test_id",
    "Platform": "platform",
    "Module": "module",
    "Test Scenario": "scenario",
    "Execution Steps": "steps",
    "Primary Tool(s)": "tools",
    "Execution Type": "execution_type",
    "OWASP / Standard": "standard",
    "Default Severity": "severity",
    "Expected Result": "expected",
    "Status": "status",
    "Actual Result": "actual",
    "HTTP/Result Code": "result_code",
    "Evidence Link/Ref": "evidence",
    "Finding ID": "finding_id",
    "Tester": "tester",
    "Retest Status": "retest",
}


@dataclass
class TestCase:
    test_id: str
    platform: str = ""
    module: str = ""
    scenario: str = ""
    steps: str = ""
    tools: str = ""
    execution_type: str = ""      # Automated | Manual | Hybrid
    standard: str = ""            # OWASP / ASVS mapping
    severity: str = ""
    expected: str = ""
    status: str = "Not Run"
    actual: str = ""
    result_code: str = ""
    evidence: str = ""
    finding_id: str = ""
    tester: str = ""
    retest: str = ""
    row: int = 0                  # 1-based worksheet row, for write-back

    @property
    def is_automated(self) -> bool:
        return self.execution_type.strip().lower() in ("automated", "hybrid")


def load_test_cases(path: str, sheet: str = TEST_CASE_SHEET) -> list[TestCase]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet {sheet!r} not found in {path}; sheets={wb.sheetnames}")
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=False)
    header_cells = next(rows)
    headers = {}
    for idx, cell in enumerate(header_cells):
        label = (cell.value or "").strip() if isinstance(cell.value, str) else cell.value
        if label in _COLUMN_MAP:
            headers[_COLUMN_MAP[label]] = idx

    if "test_id" not in headers:
        raise ValueError("could not locate 'Test ID' column in test plan")

    cases: list[TestCase] = []
    for r_index, row in enumerate(rows, start=2):
        values = [c.value for c in row]
        test_id = values[headers["test_id"]] if headers["test_id"] < len(values) else None
        if test_id in (None, ""):
            continue
        kwargs = {}
        for attr, col in headers.items():
            v = values[col] if col < len(values) else None
            kwargs[attr] = "" if v is None else str(v)
        kwargs["row"] = r_index
        cases.append(TestCase(**kwargs))
    wb.close()
    return cases


def group_by_module(cases: list[TestCase]) -> dict[str, list[TestCase]]:
    out: dict[str, list[TestCase]] = {}
    for c in cases:
        out.setdefault(c.module or "Unspecified", []).append(c)
    return out
