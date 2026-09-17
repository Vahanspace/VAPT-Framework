import os

import openpyxl
import pytest

from vaptframework.reporting.xlsx_writer import write_findings, REGISTER_SHEET, DASHBOARD_SHEET
from vaptframework.core.findings import Finding, FindingsRegister
from vaptframework.core.severity import Severity


def _wb(path):
    wb = openpyxl.Workbook()
    wb.active.title = DASHBOARD_SHEET
    dash = wb[DASHBOARD_SHEET]
    dash["D4"] = "Open Findings"; dash["E4"] = "Count"
    for i, lbl in enumerate(["Critical", "High", "Medium", "Low", "Informational"]):
        dash.cell(row=5 + i, column=4, value=lbl)
        dash.cell(row=5 + i, column=5, value=0)
    dash.cell(row=10, column=4, value="Total Open"); dash.cell(row=10, column=5, value=0)
    reg = wb.create_sheet(REGISTER_SHEET)
    reg.append(["Finding ID", "Test ID", "Title", "Platform", "Module", "Severity", "CVSS",
                "Status", "Description", "Evidence", "Remediation", "Owner", "Target Date", "Retest Result"])
    reg.append(["OLD", "", "stale row", "", "", "", "", "", "", "", ""])  # should be cleared
    wb.save(path)


def test_write_findings_populates_register_and_dashboard(tmp_path):
    p = os.path.join(tmp_path, "plan.xlsx")
    _wb(p)
    reg = FindingsRegister()
    reg.add(Finding(title="Crit", severity=Severity.CRITICAL, module="Auth", location="a"))
    reg.add(Finding(title="Crit2", severity=Severity.CRITICAL, module="Auth", location="b"))
    reg.add(Finding(title="Low", severity=Severity.LOW, module="Config", location="c"))

    out = write_findings(p, reg)
    wb = openpyxl.load_workbook(out)
    rs = wb[REGISTER_SHEET]
    titles = [rs.cell(row=r, column=3).value for r in range(2, rs.max_row + 1)]
    assert "stale row" not in titles          # cleared
    assert "Crit" in titles and "Low" in titles
    dash = wb[DASHBOARD_SHEET]
    assert dash.cell(row=5, column=5).value == 2   # Critical
    assert dash.cell(row=8, column=5).value == 1   # Low
    assert dash.cell(row=10, column=5).value == 3  # Total open
